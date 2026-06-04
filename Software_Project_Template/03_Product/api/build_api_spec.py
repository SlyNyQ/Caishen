"""Build the Caishen API specification DOCX."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(__file__).with_name("api_specification.docx")
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(33, 39, 43)
MUTED = RGBColor(95, 101, 106)
LIGHT_GRAY = "F2F4F7"
LIGHT_GOLD = "FFF7DF"
GOLD = RGBColor(183, 137, 22)


def set_font(run, *, size: float, color: RGBColor = INK, bold: bool = False, italic: bool = False) -> None:
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold
    run.italic = italic


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    settings = {
        "Title": (25, GOLD, 0, 6),
        "Subtitle": (13, MUTED, 0, 16),
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in settings.items():
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = name.startswith("Heading")

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, *, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = tc_mar.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            tc_mar.append(tag)
        tag.set(qn("w:w"), str(value))
        tag.set(qn("w:type"), "dxa")


def set_paragraph_shading(paragraph, fill: str) -> None:
    paragraph_pr = paragraph._p.get_or_add_pPr()
    shading = paragraph_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        paragraph_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def mark_header_row(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_table_geometry(table, widths: list[int]) -> None:
    table.autofit = False
    table_pr = table._tbl.tblPr
    layout = table_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        table_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    table_width = table_pr.find(qn("w:tblW"))
    table_width.set(qn("w:w"), str(sum(widths)))
    table_width.set(qn("w:type"), "dxa")
    indent = table_pr.find(qn("w:tblInd"))
    if indent is None:
        indent = OxmlElement("w:tblInd")
        table_pr.append(indent)
    indent.set(qn("w:w"), "120")
    indent.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    header = table.rows[0].cells
    for index, label in enumerate(headers):
        set_cell_shading(header[index], LIGHT_GRAY)
        paragraph = header[index].paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        set_font(paragraph.add_run(label), size=9.5, color=DARK_BLUE, bold=True)
    mark_header_row(table.rows[0])
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            if row_index % 2:
                set_cell_shading(cells[index], "FAFBFC")
            paragraph = cells[index].paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            set_font(paragraph.add_run(value), size=9.2)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_bullet(doc: Document, text: str) -> None:
    doc.add_paragraph(text, style="List Bullet")


def add_code(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.12)
    paragraph.paragraph_format.right_indent = Inches(0.12)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    set_paragraph_shading(paragraph, "F7F7F5")
    run = paragraph.add_run(text)
    run.font.name = "Consolas"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Consolas")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Consolas")
    run.font.size = Pt(8.5)


def add_callout(doc: Document, label: str, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.12)
    paragraph.paragraph_format.right_indent = Inches(0.12)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    set_paragraph_shading(paragraph, LIGHT_GOLD)
    set_font(paragraph.add_run(f"{label}: "), size=10, color=GOLD, bold=True)
    set_font(paragraph.add_run(text), size=10)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run("Page ")
    set_font(run, size=9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instruction)
    run._r.append(end)


def configure_page(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(header.add_run("CAISHEN API  |  VERSION 1.0"), size=8.5, color=MUTED, bold=True)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_field(footer)


def build() -> None:
    doc = Document()
    configure_page(doc)
    configure_styles(doc)
    core = doc.core_properties
    core.title = "Caishen API Specification"
    core.subject = "Investment analysis API contracts and safety boundaries"
    core.author = "Caishen"
    core.keywords = "Caishen, API, FastAPI, investment analysis"

    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_after = Pt(4)
    set_font(kicker.add_run("PRODUCT API SPECIFICATION"), size=9, color=GOLD, bold=True)
    doc.add_paragraph("Caishen Investment Analysis API", style="Title")
    doc.add_paragraph(
        "Plain-English market research, public-source context, and educational strategy scenarios.",
        style="Subtitle",
    )
    for label, value in (
        ("Status", "Implementation-aligned"),
        ("Version", "1.0"),
        ("Runtime", "FastAPI / AWS Lambda / API Gateway"),
        ("Last reviewed", "June 4, 2026"),
    ):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        set_font(paragraph.add_run(f"{label}: "), size=10, color=DARK_BLUE, bold=True)
        set_font(paragraph.add_run(value), size=10)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    add_callout(
        doc,
        "Safety boundary",
        "Caishen is read-only and educational. It never executes trades, guarantees outcomes, or issues personalized buy/sell instructions.",
    )

    doc.add_heading("1. Purpose And Scope", level=1)
    doc.add_paragraph(
        "The Caishen API accepts natural-language investment research questions plus optional structured context. "
        "It returns a plain-English answer and a structured analysis brief for the frontend insights panel."
    )
    add_bullet(doc, "Analyze and compare public-market tickers using yfinance-backed snapshots.")
    add_bullet(doc, "Summarize user-supplied public HTTP/HTTPS sources after guarded retrieval.")
    add_bullet(doc, "Retrieve curated educational notes for investment concepts and risk framing.")
    add_bullet(doc, "Describe hypothetical strategy scenarios without making a personalized recommendation.")

    doc.add_heading("2. Service Contract", level=1)
    add_table(
        doc,
        ["Method", "Path", "Purpose", "Success"],
        [
            ["GET", "/health", "Liveness and application identity", "200"],
            ["GET", "/ready", "Provider and runtime readiness", "200"],
            ["POST", "/chat", "Run a Caishen research workflow", "200"],
            ["GET", "/docs", "FastAPI interactive OpenAPI reference", "200"],
        ],
        [1050, 1600, 5210, 1500],
    )
    doc.add_heading("Provider Selection", level=2)
    doc.add_paragraph(
        "A request may select mock, OpenAI, Anthropic, Google, or Bedrock. The optional model must match the "
        "server-side allowlisted model for that provider. Credentials are never accepted from the client."
    )
    add_table(
        doc,
        ["Provider", "Default model setting", "Credential path"],
        [
            ["mock", "caishen-mock", "None; deterministic offline behavior"],
            ["openai", "OPENAI_MODEL", "OPENAI_API_KEY on server"],
            ["anthropic", "ANTHROPIC_MODEL", "ANTHROPIC_API_KEY on server"],
            ["google", "GOOGLE_MODEL", "GOOGLE_API_KEY on server"],
            ["bedrock", "BEDROCK_MODEL_ID", "Lambda role / AWS credentials"],
        ],
        [1500, 2800, 5060],
    )

    doc.add_heading("3. POST /chat Request", level=1)
    add_table(
        doc,
        ["Field", "Type", "Required", "Rules"],
        [
            ["message", "string", "Yes", "1-4,000 characters"],
            ["session_id", "string | null", "No", "Reuses a client session identifier"],
            ["client_request_id", "string | null", "No", "Optional idempotency/correlation identifier"],
            ["provider", "enum | null", "No", "mock, openai, anthropic, google, bedrock"],
            ["model", "string | null", "No", "Must equal provider allowlist entry"],
            ["analysis_context", "object", "No", "Defaults to an empty context"],
        ],
        [2100, 1800, 1250, 4210],
    )
    doc.add_heading("Analysis Context", level=2)
    add_table(
        doc,
        ["Field", "Type", "Limit / values"],
        [
            ["tickers", "string[]", "Up to 8; normalized uppercase"],
            ["source_urls", "string[]", "Up to 3 public HTTP/HTTPS URLs"],
            ["risk_tolerance", "enum", "unspecified, conservative, balanced, growth"],
            ["horizon", "enum", "unspecified, short, medium, long"],
            ["goal", "string | null", "Up to 300 characters"],
        ],
        [2300, 1800, 5260],
    )
    add_code(
        doc,
        '{\n'
        '  "message": "Compare AAPL and MSFT for a long-term goal",\n'
        '  "provider": "mock",\n'
        '  "analysis_context": {\n'
        '    "tickers": ["AAPL", "MSFT"],\n'
        '    "source_urls": [],\n'
        '    "risk_tolerance": "balanced",\n'
        '    "horizon": "long",\n'
        '    "goal": "Evaluate durable growth"\n'
        '  }\n'
        '}',
    )

    doc.add_heading("4. POST /chat Response", level=1)
    add_table(
        doc,
        ["Field", "Type", "Purpose"],
        [
            ["request_id", "string", "Server request correlation identifier"],
            ["session_id", "string", "Session identifier"],
            ["route", "string", "Selected workflow route"],
            ["provider / model", "string", "Actual generation adapter and allowlisted model"],
            ["answer", "string", "Plain-English response"],
            ["analysis", "object", "Structured insights for the frontend"],
            ["tool_traces", "array", "Deterministic tool inputs, results, and warnings"],
            ["citations", "array", "Public-source or curated-note evidence"],
            ["refusal", "object | null", "Safety category and explanation"],
            ["debug", "object", "Optional server-controlled diagnostics"],
        ],
        [2200, 1800, 5360],
    )
    doc.add_heading("Structured Analysis", level=2)
    add_table(
        doc,
        ["Collection", "Content"],
        [
            ["market_snapshots", "Ticker, company, price, returns, volatility, valuation metadata, timestamp, and status"],
            ["key_takeaways", "Short evidence-grounded observations"],
            ["risk_notes", "Data, volatility, concentration, and safety caveats"],
            ["source_summaries", "URL, title, sanitized summary, fetch timestamp, and status"],
            ["strategy_scenarios", "Educational hypothetical scenarios based on stated context"],
        ],
        [2600, 6760],
    )
    add_callout(
        doc,
        "Unavailable data",
        "When market or source data cannot be verified, Caishen returns an explicit unavailable/stale status and does not invent values.",
    )

    doc.add_heading("5. Routing And Safety", level=1)
    add_table(
        doc,
        ["Route", "Typical trigger", "Behavior"],
        [
            ["market_analysis", "Ticker lookup or comparison", "Market tools plus structured analysis"],
            ["source_analysis", "One or more public source URLs", "Guarded fetch, summary, citation"],
            ["research_explanation", "Investment concept or method", "Curated RAG notes and citations"],
            ["strategy_scenario", "Educational horizon/risk scenario", "Hypothetical tradeoff framing"],
            ["clarification_route", "Insufficient context", "Asks for a more specific research question"],
            ["refusal_route", "Execution, guarantee, manipulation, or personalized directive", "Refuses and offers safe educational alternatives"],
        ],
        [2200, 2700, 4460],
    )
    doc.add_heading("Public URL Guardrails", level=2)
    add_bullet(doc, "Only HTTP and HTTPS URLs without embedded credentials are accepted.")
    add_bullet(doc, "Loopback, private, link-local, multicast, reserved, and local hostnames are blocked.")
    add_bullet(doc, "Redirect count, request time, response size, and content type are limited.")
    add_bullet(doc, "Scripts, styles, forms, navigation, and other non-content elements are removed.")
    add_bullet(doc, "Every source summary carries an explicit fetch timestamp and status.")

    doc.add_heading("6. Error Contract", level=1)
    doc.add_paragraph(
        "Application errors are translated to JSON HTTP errors. Client validation failures use FastAPI/Pydantic's "
        "422 response. Provider readiness and model invocation errors remain distinct so the frontend can present useful states."
    )
    add_table(
        doc,
        ["Status", "Code", "Meaning"],
        [
            ["403", "policy_violation / access_denied", "A guarded action or resource was denied"],
            ["422", "validation_error", "Provider/model or request validation failed"],
            ["502", "model_invocation_error", "Configured provider generation failed"],
            ["503", "resource_not_ready", "Requested provider is not configured"],
            ["500", "internal_error", "Unexpected server failure"],
        ],
        [1350, 3200, 4810],
    )
    add_code(
        doc,
        '{\n'
        '  "detail": {\n'
        '    "code": "resource_not_ready",\n'
        '    "message": "The openai provider is not configured on the server.",\n'
        '    "detail": null\n'
        '  }\n'
        '}',
    )

    doc.add_heading("7. Deployment Notes", level=1)
    add_bullet(doc, "CloudFront routes /api/* to API Gateway and serves the static Next.js export from private S3.")
    add_bullet(doc, "The Lambda image builds local RAG chunks from curated Markdown during image creation.")
    add_bullet(doc, "Provider keys and AWS credentials remain outside frontend bundles and chat payloads.")
    add_bullet(doc, "Terraform apply/destroy requires explicit operator review, especially when migrating older resource names.")
    add_bullet(doc, "Health, readiness, schema, refusal, and unavailable-data behavior must be smoke-tested after deployment.")

    doc.save(OUTPUT)


if __name__ == "__main__":
    build()
